"""Relatórios gerenciais: faturamento no tempo, itens, mecânicos e clientes.

Como no painel, só contam as OS FINALIZADAS, pela data da OS. As funções
devolvem dataclasses simples (valores em centavos); a tela de Relatórios
desenha os gráficos e a exportação grava as mesmas tabelas em Excel ou CSV.
"""

import csv
import sqlite3
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

from oficina.banco import para_data
from oficina.dinheiro import formatar_reais
from oficina.modelos import STATUS_CANCELADA, STATUS_FINALIZADA, TIPOS_ITEM
from oficina.periodos import MESES, deslocar_mes, formatar_data, nome_mes, semana
from oficina.servicos import painel

DIA, SEMANA, MES = "dia", "semana", "mes"
DIAS_DA_SEMANA = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado", "Domingo"]


@dataclass
class Fatia:
    """Um dia, uma semana ou um mês do gráfico de faturamento."""
    inicio: date
    fim: date
    rotulo: str  # curto, para o eixo: "out/26", "06/10"
    titulo: str  # completo, para a dica: "Outubro de 2026"
    quantidade_os: int = 0
    pecas: int = 0
    mao_de_obra: int = 0
    terceiros: int = 0
    descontos: int = 0
    total: int = 0

    @property
    def bruto(self) -> int:
        """Soma dos itens, antes dos descontos (a altura da coluna no gráfico)."""
        return self.pecas + self.mao_de_obra + self.terceiros

    @property
    def ticket_medio(self) -> int:
        return self.total // self.quantidade_os if self.quantidade_os else 0


@dataclass
class ItemVendido:
    descricao: str
    tipo: str
    quantidade: float
    vezes: int  # em quantas OS apareceu
    total: int


@dataclass
class DesempenhoMecanico:
    mecanico_id: int | None
    nome: str
    quantidade_os: int
    faturamento: int
    mao_de_obra: int
    comissao: int

    @property
    def ticket_medio(self) -> int:
        return self.faturamento // self.quantidade_os if self.quantidade_os else 0


@dataclass
class Clientes:
    atendidos: int
    novos: int  # primeira OS da história do cliente foi neste período
    voltaram: int  # mais de uma OS dentro do período

    @property
    def ja_eram_clientes(self) -> int:
        return self.atendidos - self.novos


@dataclass
class DiaDaSemana:
    nome: str
    quantidade_os: int
    faturamento: int


@dataclass
class Conversao:
    orcamentos: int
    aprovados: int

    @property
    def taxa(self) -> float | None:
        """Fração dos orçamentos que viraram OS (None se não houve orçamento)."""
        return self.aprovados / self.orcamentos if self.orcamentos else None


# ---------------------------------------------------------------- faturamento no tempo

def agrupamento_para(inicio: date, fim: date) -> str:
    """Por dia em até duas semanas, por semana em até três meses, senão por mês."""
    dias = (fim - inicio).days + 1
    if dias <= 14:
        return DIA
    if dias <= 93:
        return SEMANA
    return MES


def _inicio_da_fatia(dia: date, agrupamento: str) -> date:
    if agrupamento == DIA:
        return dia
    if agrupamento == SEMANA:
        return semana(dia)[0]
    return dia.replace(day=1)


def _proxima(inicio: date, agrupamento: str) -> date:
    if agrupamento == DIA:
        return inicio + timedelta(days=1)
    if agrupamento == SEMANA:
        return inicio + timedelta(days=7)
    return deslocar_mes(inicio, 1)


def _nova_fatia(chave: date, agrupamento: str, limite_inicio: date, limite_fim: date) -> Fatia:
    """Fatia que começa em `chave`, cortada nos limites do período (ex.: a 1ª semana de um mês)."""
    if agrupamento == DIA:
        titulo = f"{DIAS_DA_SEMANA[chave.weekday()]}, {formatar_data(chave)}"
        return Fatia(chave, chave, chave.strftime("%d/%m"), titulo)
    if agrupamento == SEMANA:
        inicio, fim = max(chave, limite_inicio), min(chave + timedelta(days=6), limite_fim)
        return Fatia(inicio, fim, inicio.strftime("%d/%m"), f"Semana de {inicio:%d/%m} a {formatar_data(fim)}")
    inicio, fim = max(chave, limite_inicio), min(deslocar_mes(chave, 1) - timedelta(days=1), limite_fim)
    return Fatia(inicio, fim, f"{MESES[chave.month - 1][:3]}/{chave:%y}", nome_mes(chave).capitalize())


def faturamento_por_periodo(conn: sqlite3.Connection, inicio: date, fim: date,
                            agrupamento: str | None = None) -> list[Fatia]:
    """Faturamento por dia, semana ou mês, incluindo os períodos sem movimento.

    Sábados e domingos sem OS ficam de fora do agrupamento por dia (a oficina
    não abre), para o gráfico não ter buracos que não significam nada.
    """
    agrupamento = agrupamento or agrupamento_para(inicio, fim)
    fatias: dict[date, Fatia] = {}
    atual = _inicio_da_fatia(inicio, agrupamento)
    while atual <= fim:
        if agrupamento != DIA or atual.weekday() < 5:
            fatias[atual] = _nova_fatia(atual, agrupamento, inicio, fim)
        atual = _proxima(atual, agrupamento)

    for linha in conn.execute(
            """SELECT data, total_pecas, total_mao_de_obra, total_terceiros, desconto, total
               FROM ordens_servico WHERE status = ? AND data BETWEEN ? AND ?""",
            (STATUS_FINALIZADA, inicio.isoformat(), fim.isoformat())):
        chave = _inicio_da_fatia(para_data(linha["data"]), agrupamento)
        if chave not in fatias:  # OS num sábado ou domingo
            fatias[chave] = _nova_fatia(chave, agrupamento, inicio, fim)
        fatia = fatias[chave]
        fatia.quantidade_os += 1
        fatia.pecas += linha["total_pecas"]
        fatia.mao_de_obra += linha["total_mao_de_obra"]
        fatia.terceiros += linha["total_terceiros"]
        fatia.descontos += linha["desconto"]
        fatia.total += linha["total"]
    return [fatias[chave] for chave in sorted(fatias)]


# ---------------------------------------------------------------- itens, mecânicos, clientes

def itens_mais_vendidos(conn: sqlite3.Connection, inicio: date, fim: date, limite: int | None = 10,
                        tipo: str | None = None) -> list[ItemVendido]:
    """Itens que mais faturaram. Descrições iguais a menos de acento, maiúsculas e espaços são somadas juntas."""
    sql = """
        SELECT TRIM(MIN(i.descricao)) AS descricao, i.tipo, SUM(i.quantidade) AS quantidade,
               COUNT(DISTINCT i.os_id) AS vezes, SUM(i.total) AS total
        FROM itens_os i JOIN ordens_servico o ON o.id = i.os_id
        WHERE o.status = ? AND o.data BETWEEN ? AND ?
    """
    parametros: list = [STATUS_FINALIZADA, inicio.isoformat(), fim.isoformat()]
    if tipo:
        sql += " AND i.tipo = ?"
        parametros.append(tipo)
    sql += " GROUP BY normalizar(TRIM(i.descricao)), i.tipo ORDER BY total DESC, descricao"
    if limite:
        sql += f" LIMIT {int(limite)}"
    return [ItemVendido(l["descricao"], l["tipo"], l["quantidade"], l["vezes"], l["total"])
            for l in conn.execute(sql, parametros)]


def desempenho_mecanicos(conn: sqlite3.Connection, inicio: date, fim: date) -> list[DesempenhoMecanico]:
    """OS, faturamento, mão de obra e comissão de cada mecânico (o nome é o do cadastro atual)."""
    linhas = conn.execute(
        """SELECT o.mecanico_id, COALESCE(m.nome, MAX(o.mecanico_nome), '') AS nome, COUNT(*) AS quantidade,
                  SUM(o.total) AS faturamento, SUM(o.total_mao_de_obra) AS mao, SUM(o.comissao) AS comissao
           FROM ordens_servico o LEFT JOIN mecanicos m ON m.id = o.mecanico_id
           WHERE o.status = ? AND o.data BETWEEN ? AND ?
           GROUP BY o.mecanico_id ORDER BY mao DESC, nome""",
        (STATUS_FINALIZADA, inicio.isoformat(), fim.isoformat()))
    return [DesempenhoMecanico(l["mecanico_id"], l["nome"] or "Sem mecânico", l["quantidade"], l["faturamento"],
                               l["mao"], l["comissao"]) for l in linhas]


def clientes_atendidos(conn: sqlite3.Connection, inicio: date, fim: date) -> Clientes:
    """Clientes com OS no período: quantos são novos e quantos voltaram mais de uma vez nele."""
    linha = conn.execute(
        """WITH primeira AS (
               SELECT cliente_id, MIN(data) AS data FROM ordens_servico
               WHERE status = ? AND cliente_id IS NOT NULL GROUP BY cliente_id),
           no_periodo AS (
               SELECT cliente_id, COUNT(*) AS visitas FROM ordens_servico
               WHERE status = ? AND cliente_id IS NOT NULL AND data BETWEEN ? AND ? GROUP BY cliente_id)
           SELECT COUNT(*) AS atendidos,
                  COALESCE(SUM(p.data >= ?), 0) AS novos,
                  COALESCE(SUM(n.visitas > 1), 0) AS voltaram
           FROM no_periodo n JOIN primeira p ON p.cliente_id = n.cliente_id""",
        (STATUS_FINALIZADA, STATUS_FINALIZADA, inicio.isoformat(), fim.isoformat(), inicio.isoformat()),
    ).fetchone()
    return Clientes(linha["atendidos"], linha["novos"], linha["voltaram"])


def movimento_por_dia_da_semana(conn: sqlite3.Connection, inicio: date, fim: date) -> list[DiaDaSemana]:
    """OS e faturamento por dia da semana. Sábado e domingo só aparecem se tiverem OS."""
    contagem = {indice: DiaDaSemana(nome, 0, 0) for indice, nome in enumerate(DIAS_DA_SEMANA)}
    for linha in conn.execute(
            """SELECT data, total FROM ordens_servico WHERE status = ? AND data BETWEEN ? AND ?""",
            (STATUS_FINALIZADA, inicio.isoformat(), fim.isoformat())):
        dia = contagem[para_data(linha["data"]).weekday()]
        dia.quantidade_os += 1
        dia.faturamento += linha["total"]
    return [dia for indice, dia in contagem.items() if indice < 5 or dia.quantidade_os]


def conversao_orcamentos(conn: sqlite3.Connection, inicio: date, fim: date) -> Conversao:
    """Orçamentos feitos no período e quantos deles viraram OS (não cancelada)."""
    linha = conn.execute(
        """SELECT COUNT(*) AS orcamentos,
                  COALESCE(SUM(EXISTS (SELECT 1 FROM ordens_servico o
                                       WHERE o.orcamento_id = orc.id AND o.status <> ?)), 0) AS aprovados
           FROM orcamentos orc WHERE orc.data BETWEEN ? AND ?""",
        (STATUS_CANCELADA, inicio.isoformat(), fim.isoformat()),
    ).fetchone()
    return Conversao(linha["orcamentos"], linha["aprovados"])


# ---------------------------------------------------------------- exportação

COLUNAS_OS = ["Nº", "Data", "Mecânico", "Cliente", "Placa", "Modelo", "Peças", "Mão de obra", "Terceirizados",
              "Desconto", "Total", "Comissão", "Comissão paga"]
_COLUNAS_OS_EM_REAIS = {"Peças", "Mão de obra", "Terceirizados", "Desconto", "Total", "Comissão"}


def _linhas_os(conn: sqlite3.Connection, inicio: date, fim: date) -> list[list]:
    """Uma linha por OS finalizada do período (valores em centavos, data como date)."""
    return [
        [l["numero"], para_data(l["data"]), l["mecanico_nome"], l["cliente_nome"], l["placa"], l["modelo"],
         l["total_pecas"], l["total_mao_de_obra"], l["total_terceiros"], l["desconto"], l["total"], l["comissao"],
         "Sim" if l["comissao_paga"] else ("Não" if l["comissao"] else "")]
        for l in conn.execute(
            "SELECT * FROM ordens_servico WHERE status = ? AND data BETWEEN ? AND ? ORDER BY data, numero",
            (STATUS_FINALIZADA, inicio.isoformat(), fim.isoformat()))
    ]


def exportar_csv(conn: sqlite3.Connection, inicio: date, fim: date, caminho: str | Path) -> Path:
    """Lista de OS do período em CSV no padrão do Excel brasileiro (; e vírgula decimal)."""
    caminho = Path(caminho)
    with caminho.open("w", newline="", encoding="utf-8-sig") as arquivo:
        escritor = csv.writer(arquivo, delimiter=";")
        escritor.writerow(COLUNAS_OS)
        for linha in _linhas_os(conn, inicio, fim):
            escritor.writerow([
                formatar_reais(valor, simbolo=False).replace(".", "") if nome in _COLUNAS_OS_EM_REAIS
                else formatar_data(valor) if isinstance(valor, date) else valor
                for nome, valor in zip(COLUNAS_OS, linha, strict=True)
            ])
    return caminho


def exportar_excel(conn: sqlite3.Connection, inicio: date, fim: date, caminho: str | Path) -> Path:
    """Planilha com o resumo do período, o faturamento no tempo, itens, mecânicos e a lista de OS."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    formato_reais = '"R$" #,##0.00'
    cabecalho_fonte = Font(bold=True, color="FFFFFF")
    cabecalho_fundo = PatternFill("solid", fgColor="1A365C")

    def planilha(livro, titulo: str, colunas: list[str], linhas: list[list], em_reais: set[str],
                 larguras: dict[str, int] | None = None):
        aba = livro.create_sheet(titulo)
        aba.append(colunas)
        for celula in aba[1]:
            celula.font, celula.fill = cabecalho_fonte, cabecalho_fundo
            celula.alignment = Alignment(horizontal="center", vertical="center")
        for linha in linhas:
            aba.append([valor / 100 if nome in em_reais and isinstance(valor, int) else valor
                        for nome, valor in zip(colunas, linha, strict=True)])
        for indice, nome in enumerate(colunas, start=1):
            letra = get_column_letter(indice)
            aba.column_dimensions[letra].width = (larguras or {}).get(nome, max(12, len(nome) + 4))
            for (celula,) in aba.iter_rows(min_row=2, min_col=indice, max_col=indice):
                if nome in em_reais:
                    celula.number_format = formato_reais
                elif isinstance(celula.value, date):
                    celula.number_format = "DD/MM/YYYY"
        aba.freeze_panes = "A2"
        if linhas:
            aba.auto_filter.ref = aba.dimensions
        return aba

    resumo = painel.resumo_periodo(conn, inicio, fim)
    clientes = clientes_atendidos(conn, inicio, fim)
    conversao = conversao_orcamentos(conn, inicio, fim)
    livro = Workbook()
    livro.remove(livro.active)

    aba = planilha(livro, "Resumo", ["Indicador", "Valor"], [
        ["Período", f"{formatar_data(inicio)} a {formatar_data(fim)}"],
        ["OS finalizadas", resumo.quantidade_os],
        ["Faturamento", resumo.faturamento / 100],
        ["Peças", resumo.pecas / 100],
        ["Mão de obra", resumo.mao_de_obra / 100],
        ["Terceirizados", resumo.terceiros / 100],
        ["Descontos", resumo.descontos / 100],
        ["Ticket médio", resumo.ticket_medio / 100],
        ["Comissões", resumo.comissao / 100],
        ["Clientes atendidos", clientes.atendidos],
        ["Clientes novos", clientes.novos],
        ["Clientes que voltaram mais de uma vez", clientes.voltaram],
        ["Orçamentos feitos", conversao.orcamentos],
        ["Orçamentos que viraram OS", conversao.aprovados],
    ], set(), {"Indicador": 38, "Valor": 26})
    for (celula,) in aba.iter_rows(min_row=4, max_row=10, min_col=2, max_col=2):
        celula.number_format = formato_reais

    fatias = faturamento_por_periodo(conn, inicio, fim)
    planilha(livro, "Faturamento", ["Período", "OS", "Peças", "Mão de obra", "Terceirizados", "Descontos", "Total",
                                    "Ticket médio"],
             [[f.titulo, f.quantidade_os, f.pecas, f.mao_de_obra, f.terceiros, f.descontos, f.total, f.ticket_medio]
              for f in fatias],
             {"Peças", "Mão de obra", "Terceirizados", "Descontos", "Total", "Ticket médio"}, {"Período": 30})
    planilha(livro, "Itens", ["Descrição", "Tipo", "Quantidade", "Nº de OS", "Total"],
             [[i.descricao, TIPOS_ITEM.get(i.tipo, i.tipo), i.quantidade, i.vezes, i.total]
              for i in itens_mais_vendidos(conn, inicio, fim, limite=None)],
             {"Total"}, {"Descrição": 40, "Tipo": 22})
    planilha(livro, "Mecânicos", ["Mecânico", "OS", "Faturamento", "Mão de obra", "Ticket médio", "Comissão"],
             [[m.nome, m.quantidade_os, m.faturamento, m.mao_de_obra, m.ticket_medio, m.comissao]
              for m in desempenho_mecanicos(conn, inicio, fim)],
             {"Faturamento", "Mão de obra", "Ticket médio", "Comissão"}, {"Mecânico": 28})
    planilha(livro, "OS", COLUNAS_OS, _linhas_os(conn, inicio, fim), _COLUNAS_OS_EM_REAIS,
             {"Cliente": 30, "Mecânico": 22, "Modelo": 18})

    caminho = Path(caminho)
    livro.save(caminho)
    return caminho
