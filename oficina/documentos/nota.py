"""Layout da nota de OS e do orçamento (planilha Excel convertida em PDF).

O visual é o mesmo da versão original. As únicas mudanças em relação ao
app.py antigo são: nome/endereço da oficina e data vêm por parâmetro (para
reimprimir uma OS antiga com a data dela) e a conversão para PDF fica em
oficina/documentos/pdf.py.
"""

import os
import re
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.page import PageMargins

from oficina.documentos.pdf import exportar_pdf

OFICINA_PADRAO = {"nome": "MINHA OFICINA", "endereco": ""}


def _nome_arquivo(texto):
    """Remove caracteres que o Windows não aceita em nome de arquivo."""
    return re.sub(r'[\\/:*?"<>|]', "", str(texto)).strip()

# ==========================================
# PALETA DE CORES E ESTILOS GLOBAIS
# ==========================================
COR_AZUL_ESCURO = "1A365C"    
COR_AZUL_CLARO = "F1F5F9"     
COR_LINHAS_SUAVES = "CBD5E1"  

FILL_ESCURO = PatternFill(start_color=COR_AZUL_ESCURO, end_color=COR_AZUL_ESCURO, fill_type="solid")
FILL_CLARO = PatternFill(start_color=COR_AZUL_CLARO, end_color=COR_AZUL_CLARO, fill_type="solid")
FILL_BRANCO = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

FONT_TITULO_OFICINA = Font(name="Segoe UI", bold=True, color=COR_AZUL_ESCURO, size=16)
FONT_BRANCA_BOLD = Font(name="Segoe UI", bold=True, color="FFFFFF", size=10)
FONT_BOLD_AZUL = Font(name="Segoe UI", bold=True, size=10, color=COR_AZUL_ESCURO)
FONT_NORMAL_PRETO = Font(name="Segoe UI", bold=False, size=10, color="000000")
FONT_BOLD_PRETO = Font(name="Segoe UI", bold=True, size=10, color="000000")

BORDA_SUAVE = Side(style='thin', color=COR_LINHAS_SUAVES)
BORDER_PADRAO = Border(left=BORDA_SUAVE, right=BORDA_SUAVE, top=BORDA_SUAVE, bottom=BORDA_SUAVE)
SEM_BORDA = Border(left=Side(style=None), right=Side(style=None), top=Side(style=None), bottom=Side(style=None))

CENTER_ALIGN = Alignment(horizontal='center', vertical='center')
LEFT_ALIGN = Alignment(horizontal='left', vertical='center')
RIGHT_ALIGN = Alignment(horizontal='right', vertical='center')

def aplicar_estilo(ws, range_str, fill=None, font=None, alignment=None):
    for row in ws[range_str]:
        for cell in row:
            cell.border = BORDER_PADRAO
            if fill: cell.fill = fill
            if font: cell.font = font
            if alignment: cell.alignment = alignment


# ==========================================
# FUNÇÃO: GERAR ORDEM DE SERVIÇO (OS)
# ==========================================
def gerar_os_e_pdf(numero_os, cliente, mecanico, veiculo, itens, desconto, observacoes, pasta_destino, data_emissao=None, oficina=None):
    oficina = oficina or OFICINA_PADRAO
    data_emissao = data_emissao or datetime.now()
    nome_base = f"OS_{_nome_arquivo(veiculo['placa'])}_{numero_os}"
    caminho_excel = os.path.abspath(os.path.join(pasta_destino, f"{nome_base}.xlsx"))
    caminho_pdf = os.path.abspath(os.path.join(pasta_destino, f"{nome_base}.pdf"))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"OS_{numero_os}"

    for i in range(1, 100): ws.row_dimensions[i].height = 24
    ws.column_dimensions['A'].width = 15  
    ws.column_dimensions['B'].width = 18  
    ws.column_dimensions['C'].width = 13  
    ws.column_dimensions['D'].width = 15  
    ws.column_dimensions['E'].width = 15  
    ws.column_dimensions['F'].width = 15  

    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 3   
    ws.row_dimensions[4].height = 16  

    ws.merge_cells('A1:F1')
    ws['A1'] = f"ORDEM DE SERVIÇO - {oficina['nome']}"
    ws['A1'].font = FONT_TITULO_OFICINA
    ws['A1'].alignment = Alignment(horizontal='center', vertical='bottom')

    ws.merge_cells('A2:F2')
    ws['A2'] = f"Endereço: {oficina['endereco']}"
    ws['A2'].font = FONT_NORMAL_PRETO
    ws['A2'].alignment = Alignment(horizontal='center', vertical='top')

    ws.merge_cells('A3:F3')
    for cell in ws['A3:F3'][0]:
        cell.fill = FILL_ESCURO
        cell.border = SEM_BORDA

    ws['A5'] = "OS Nº:"
    ws.merge_cells('B5:D5'); ws['B5'] = numero_os
    ws['E5'] = "Data:"
    ws['F5'] = data_emissao.strftime('%d/%m/%Y')
    
    aplicar_estilo(ws, 'A5:A5', FILL_CLARO, FONT_BOLD_AZUL, LEFT_ALIGN)
    aplicar_estilo(ws, 'B5:D5', FILL_BRANCO, FONT_NORMAL_PRETO, LEFT_ALIGN)
    aplicar_estilo(ws, 'E5:E5', FILL_CLARO, FONT_BOLD_AZUL, LEFT_ALIGN)
    aplicar_estilo(ws, 'F5:F5', FILL_BRANCO, FONT_NORMAL_PRETO, LEFT_ALIGN) 

    ws['A6'] = "Mecânico:"
    ws.merge_cells('B6:F6'); ws['B6'] = mecanico
    aplicar_estilo(ws, 'A6:A6', FILL_CLARO, FONT_BOLD_AZUL, LEFT_ALIGN)
    aplicar_estilo(ws, 'B6:F6', FILL_BRANCO, FONT_NORMAL_PRETO, LEFT_ALIGN)

    ws.row_dimensions[7].height = 10 

    ws.merge_cells('A8:C8'); ws['A8'] = "DADOS DO CLIENTE"
    ws.merge_cells('D8:F8'); ws['D8'] = "DADOS DO VEÍCULO"
    aplicar_estilo(ws, 'A8:F8', FILL_CLARO, FONT_BOLD_AZUL, CENTER_ALIGN)

    ws['A9'] = "Nome:"; ws.merge_cells('B9:C9'); ws['B9'] = cliente['nome']
    ws['D9'] = "Modelo:"; ws.merge_cells('E9:F9'); ws['E9'] = veiculo['modelo']
    ws['A10'] = "CPF/CNPJ:"; ws.merge_cells('B10:C10'); ws['B10'] = cliente['documento']
    ws['D10'] = "Placa:"; ws.merge_cells('E10:F10'); ws['E10'] = veiculo['placa']
    ws['A11'] = "Telefone:"; ws.merge_cells('B11:C11'); ws['B11'] = cliente['telefone']
    ws['D11'] = "Ano / KM:"; ws.merge_cells('E11:F11'); ws['E11'] = f"{veiculo['ano']} / {veiculo['km']}"

    for r in ['9', '10', '11']:
        aplicar_estilo(ws, f'A{r}:A{r}', FILL_CLARO, FONT_BOLD_AZUL, LEFT_ALIGN)
        aplicar_estilo(ws, f'B{r}:C{r}', FILL_BRANCO, FONT_NORMAL_PRETO, LEFT_ALIGN)
        aplicar_estilo(ws, f'D{r}:D{r}', FILL_CLARO, FONT_BOLD_AZUL, LEFT_ALIGN)
        aplicar_estilo(ws, f'E{r}:F{r}', FILL_BRANCO, FONT_NORMAL_PRETO, LEFT_ALIGN)

    ws.row_dimensions[12].height = 10 

    ws.merge_cells('A13:C13'); ws['A13'] = "Descrição de Peças e Serviços"
    ws['D13'] = "Qtd"; ws['E13'] = "Vlr. Unit"; ws['F13'] = "Total"
    aplicar_estilo(ws, 'A13:F13', FILL_CLARO, FONT_BOLD_AZUL, CENTER_ALIGN)

    row = 14
    for item in itens:
        ws.merge_cells(f'A{row}:C{row}'); ws[f'A{row}'] = item['desc']
        ws[f'D{row}'] = item['qtd']
        ws[f'E{row}'] = item['unit']
        ws[f'F{row}'] = f"=D{row}*E{row}"
        
        aplicar_estilo(ws, f'A{row}:C{row}', FILL_BRANCO, FONT_NORMAL_PRETO, LEFT_ALIGN)
        aplicar_estilo(ws, f'D{row}:D{row}', FILL_BRANCO, FONT_NORMAL_PRETO, CENTER_ALIGN)
        aplicar_estilo(ws, f'E{row}:F{row}', FILL_BRANCO, FONT_NORMAL_PRETO, RIGHT_ALIGN)

        ws[f'E{row}'].number_format = '"R$ "#,##0.00'
        ws[f'F{row}'].number_format = '"R$ "#,##0.00'
        row += 1

    ws.row_dimensions[row].height = 10
    row += 1

    tot_row = row
    ws[f'E{tot_row}'] = "Subtotal:"; ws[f'F{tot_row}'] = f"=SUM(F14:F{tot_row-2})"
    ws[f'E{tot_row+1}'] = "Desconto:"; ws[f'F{tot_row+1}'] = float(desconto)
    ws[f'E{tot_row+2}'] = "TOTAL GERAL:"; ws[f'F{tot_row+2}'] = f"=F{tot_row}-F{tot_row+1}"
    
    aplicar_estilo(ws, f'E{tot_row}:E{tot_row+1}', FILL_CLARO, FONT_BOLD_AZUL, RIGHT_ALIGN)
    aplicar_estilo(ws, f'F{tot_row}:F{tot_row+1}', FILL_BRANCO, FONT_NORMAL_PRETO, RIGHT_ALIGN)
    aplicar_estilo(ws, f'E{tot_row+2}:F{tot_row+2}', FILL_ESCURO, FONT_BRANCA_BOLD, RIGHT_ALIGN)
    for r in range(tot_row, tot_row+3): ws[f'F{r}'].number_format = '"R$ "#,##0.00'

    obs_row = tot_row + 4
    ws.merge_cells(f'A{obs_row}:F{obs_row}'); ws[f'A{obs_row}'] = "Observações:"
    aplicar_estilo(ws, f'A{obs_row}:F{obs_row}', FILL_CLARO, FONT_BOLD_AZUL, LEFT_ALIGN)
    ws.merge_cells(f'A{obs_row+1}:F{obs_row+2}'); ws[f'A{obs_row+1}'] = observacoes
    aplicar_estilo(ws, f'A{obs_row+1}:F{obs_row+2}', FILL_BRANCO, FONT_NORMAL_PRETO, Alignment(vertical='top', wrap_text=True))

    sig_row = obs_row + 5 
    ws.merge_cells(f'A{sig_row}:C{sig_row}'); ws.merge_cells(f'A{sig_row+1}:C{sig_row+1}')
    ws.merge_cells(f'D{sig_row}:F{sig_row}'); ws.merge_cells(f'D{sig_row+1}:F{sig_row+1}')
    
    ws[f'A{sig_row}'] = "________________________________________________"
    ws[f'D{sig_row}'] = "________________________________________________"
    ws[f'A{sig_row+1}'] = "Assinatura do Técnico"
    ws[f'D{sig_row+1}'] = "Assinatura do Cliente"

    aplicar_estilo(ws, f'A{sig_row}:C{sig_row+1}', None, FONT_BOLD_AZUL, CENTER_ALIGN)
    aplicar_estilo(ws, f'D{sig_row}:F{sig_row+1}', None, FONT_BOLD_AZUL, CENTER_ALIGN)
    for r in [sig_row, sig_row+1]:
        for c in ['A','B','C','D','E','F']: ws[f'{c}{r}'].border = SEM_BORDA

    ws.page_margins = PageMargins(left=0.2, right=0.2, top=0.5, bottom=0.5)
    ws.print_options.horizontalCentered = True  
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.fitToPage, ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = True, 1, 0 

    wb.save(caminho_excel)
    return exportar_pdf(caminho_excel, caminho_pdf)

# ==========================================
# FUNÇÃO: GERAR ORÇAMENTO (SIMPLIFICADO)
# ==========================================
def gerar_orcamento_e_pdf(numero_orc, veiculo, itens, desconto, observacoes, pasta_destino, data_emissao=None, oficina=None):
    oficina = oficina or OFICINA_PADRAO
    data_emissao = data_emissao or datetime.now()
    nome_base = f"ORC_{_nome_arquivo(veiculo['placa'])}_{numero_orc}"
    caminho_excel = os.path.abspath(os.path.join(pasta_destino, f"{nome_base}.xlsx"))
    caminho_pdf = os.path.abspath(os.path.join(pasta_destino, f"{nome_base}.pdf"))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"ORC_{numero_orc}"

    for i in range(1, 100): ws.row_dimensions[i].height = 24
    ws.column_dimensions['A'].width = 15; ws.column_dimensions['B'].width = 18; ws.column_dimensions['C'].width = 13  
    ws.column_dimensions['D'].width = 15; ws.column_dimensions['E'].width = 15; ws.column_dimensions['F'].width = 15  

    ws.row_dimensions[1].height = 36; ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 3; ws.row_dimensions[4].height = 16  

    ws.merge_cells('A1:F1')
    ws['A1'] = f"ORÇAMENTO - {oficina['nome']}" 
    ws['A1'].font = FONT_TITULO_OFICINA
    ws['A1'].alignment = Alignment(horizontal='center', vertical='bottom')

    ws.merge_cells('A2:F2')
    ws['A2'] = f"Endereço: {oficina['endereco']}"
    ws['A2'].font = FONT_NORMAL_PRETO
    ws['A2'].alignment = Alignment(horizontal='center', vertical='top')

    ws.merge_cells('A3:F3')
    for cell in ws['A3:F3'][0]:
        cell.fill = FILL_ESCURO
        cell.border = SEM_BORDA

    ws['A5'] = "Orçamento Nº:"
    ws.merge_cells('B5:D5'); ws['B5'] = numero_orc
    ws['E5'] = "Data:"
    ws['F5'] = data_emissao.strftime('%d/%m/%Y')
    
    aplicar_estilo(ws, 'A5:A5', FILL_CLARO, FONT_BOLD_AZUL, LEFT_ALIGN)
    aplicar_estilo(ws, 'B5:D5', FILL_BRANCO, FONT_NORMAL_PRETO, LEFT_ALIGN)
    aplicar_estilo(ws, 'E5:E5', FILL_CLARO, FONT_BOLD_AZUL, LEFT_ALIGN)
    aplicar_estilo(ws, 'F5:F5', FILL_BRANCO, FONT_NORMAL_PRETO, LEFT_ALIGN) 

    ws.row_dimensions[6].height = 10 

    ws.merge_cells('A7:F7'); ws['A7'] = "DADOS DO VEÍCULO"
    aplicar_estilo(ws, 'A7:F7', FILL_CLARO, FONT_BOLD_AZUL, CENTER_ALIGN)

    ws['A8'] = "Modelo:"
    ws.merge_cells('B8:C8'); ws['B8'] = veiculo['modelo']
    ws['D8'] = "Placa:"
    ws.merge_cells('E8:F8'); ws['E8'] = veiculo['placa']

    aplicar_estilo(ws, 'A8:A8', FILL_CLARO, FONT_BOLD_AZUL, LEFT_ALIGN)
    aplicar_estilo(ws, 'B8:C8', FILL_BRANCO, FONT_NORMAL_PRETO, LEFT_ALIGN)
    aplicar_estilo(ws, 'D8:D8', FILL_CLARO, FONT_BOLD_AZUL, LEFT_ALIGN)
    aplicar_estilo(ws, 'E8:F8', FILL_BRANCO, FONT_NORMAL_PRETO, LEFT_ALIGN)

    ws.row_dimensions[9].height = 10 

    ws.merge_cells('A10:C10'); ws['A10'] = "Descrição de Peças e Serviços"
    ws['D10'] = "Qtd"; ws['E10'] = "Vlr. Unit"; ws['F10'] = "Total"
    aplicar_estilo(ws, 'A10:F10', FILL_CLARO, FONT_BOLD_AZUL, CENTER_ALIGN)

    row = 11
    for item in itens:
        ws.merge_cells(f'A{row}:C{row}'); ws[f'A{row}'] = item['desc']
        ws[f'D{row}'] = item['qtd']
        ws[f'E{row}'] = item['unit']
        ws[f'F{row}'] = f"=D{row}*E{row}"
        
        aplicar_estilo(ws, f'A{row}:C{row}', FILL_BRANCO, FONT_NORMAL_PRETO, LEFT_ALIGN)
        aplicar_estilo(ws, f'D{row}:D{row}', FILL_BRANCO, FONT_NORMAL_PRETO, CENTER_ALIGN)
        aplicar_estilo(ws, f'E{row}:F{row}', FILL_BRANCO, FONT_NORMAL_PRETO, RIGHT_ALIGN)

        ws[f'E{row}'].number_format = '"R$ "#,##0.00'
        ws[f'F{row}'].number_format = '"R$ "#,##0.00'
        row += 1

    ws.row_dimensions[row].height = 10
    row += 1

    tot_row = row
    ws[f'E{tot_row}'] = "Subtotal:"; ws[f'F{tot_row}'] = f"=SUM(F11:F{tot_row-2})"
    ws[f'E{tot_row+1}'] = "Desconto:"; ws[f'F{tot_row+1}'] = float(desconto)
    ws[f'E{tot_row+2}'] = "TOTAL GERAL:"; ws[f'F{tot_row+2}'] = f"=F{tot_row}-F{tot_row+1}"
    
    aplicar_estilo(ws, f'E{tot_row}:E{tot_row+1}', FILL_CLARO, FONT_BOLD_AZUL, RIGHT_ALIGN)
    aplicar_estilo(ws, f'F{tot_row}:F{tot_row+1}', FILL_BRANCO, FONT_NORMAL_PRETO, RIGHT_ALIGN)
    aplicar_estilo(ws, f'E{tot_row+2}:F{tot_row+2}', FILL_ESCURO, FONT_BRANCA_BOLD, RIGHT_ALIGN)
    for r in range(tot_row, tot_row+3): ws[f'F{r}'].number_format = '"R$ "#,##0.00'

    obs_row = tot_row + 4
    ws.merge_cells(f'A{obs_row}:F{obs_row}'); ws[f'A{obs_row}'] = "Observações / Validade:"
    aplicar_estilo(ws, f'A{obs_row}:F{obs_row}', FILL_CLARO, FONT_BOLD_AZUL, LEFT_ALIGN)
    ws.merge_cells(f'A{obs_row+1}:F{obs_row+2}'); ws[f'A{obs_row+1}'] = observacoes
    aplicar_estilo(ws, f'A{obs_row+1}:F{obs_row+2}', FILL_BRANCO, FONT_NORMAL_PRETO, Alignment(vertical='top', wrap_text=True))

    ws.page_margins = PageMargins(left=0.2, right=0.2, top=0.5, bottom=0.5)
    ws.print_options.horizontalCentered = True  
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.fitToPage, ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = True, 1, 0 

    wb.save(caminho_excel)
    return exportar_pdf(caminho_excel, caminho_pdf)
